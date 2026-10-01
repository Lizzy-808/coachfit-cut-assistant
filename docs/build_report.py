"""Export docs/tradeoff_analysis.md to PDF (via headless Chrome) and Word.

    .venv/bin/python docs/build_report.py

Handles the Markdown this report uses: headings, paragraphs, bullet lists, tables,
**bold**, *italic*, `code`, and a trailing footnote after ---.
"""
import html
import re
import subprocess
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
SRC = HERE / "tradeoff_analysis.md"
STEM = "PE6201_Tradeoff_Analysis_LIU_ZEYUAN"
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"


def blocks(md: str):
    """Yield (kind, payload): h1/h2/p/ul/table/hr."""
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        s = lines[i].rstrip()
        if not s:
            i += 1
        elif s.startswith("# "):
            yield "h1", s[2:]; i += 1
        elif s.startswith("## "):
            yield "h2", s[3:]; i += 1
        elif s.startswith("---"):
            yield "hr", None; i += 1
        elif s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not set("".join(cells)) <= set("-: "):
                    rows.append(cells)
                i += 1
            yield "table", rows
        elif s.startswith("- "):
            items = []
            while i < len(lines) and lines[i].startswith("- "):
                items.append(lines[i][2:].strip()); i += 1
            yield "ul", items
        else:
            para = [s]
            i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r"(#|\||- |---)", lines[i]):
                para.append(lines[i].strip()); i += 1
            yield "p", " ".join(para)


INLINE = re.compile(r"(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|(?<![\\\w])\*[^*]+\*|`[^`]+`)")
LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def inline_html(t: str) -> str:
    out = []
    for part in INLINE.split(t.replace("\\*", "\x00")):
        if not part:
            continue
        esc = html.escape(part.replace("\x00", "*"))
        if LINK.fullmatch(part):
            text, url = LINK.fullmatch(part).groups()
            out.append(f'<a href="{html.escape(url)}">{html.escape(text)}</a>')
        elif part.startswith("**"):
            out.append(f"<strong>{html.escape(part[2:-2])}</strong>")
        elif part.startswith("`"):
            out.append(f"<code>{html.escape(part[1:-1])}</code>")
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            out.append(f"<em>{html.escape(part[1:-1])}</em>")
        else:
            out.append(esc)
    return "".join(out)


CSS = """
@page { size: A4; margin: 18mm 19mm 18mm 19mm; }
body { font-family: "Helvetica Neue", Helvetica, Arial, sans-serif; font-size: 10pt;
       line-height: 1.42; color: #1a1a1a; }
h1 { font-size: 16pt; margin: 0 0 2pt; color: #12324f; }
.byline { color: #555; margin: 0 0 12pt; font-size: 9.5pt; }
h2 { font-size: 11.5pt; margin: 13pt 0 4pt; color: #12324f; border-bottom: 0.6pt solid #c9d3dd;
     padding-bottom: 2pt; page-break-after: avoid; }
p { margin: 0 0 6pt; text-align: justify; hyphens: auto; }
ul { margin: 0 0 6pt 0; padding-left: 14pt; }
li { margin-bottom: 3pt; text-align: justify; }
table { border-collapse: collapse; width: 100%; margin: 4pt 0 6pt; font-size: 9pt;
        page-break-inside: avoid; }
th, td { border: 0.5pt solid #b8c2cc; padding: 3pt 5pt; vertical-align: top; text-align: left; }
th { background: #eef2f6; }
a { color: #1f5fa8; text-decoration: none; }
code { font-family: Menlo, Consolas, monospace; font-size: 8.6pt; }
.note { font-size: 8.4pt; color: #444; }
hr { border: 0; border-top: 0.6pt solid #c9d3dd; margin: 10pt 0 5pt; }
.foot { font-size: 8.4pt; color: #555; }
"""


def to_html(md: str) -> str:
    body, after_hr, first_p = [], False, True
    for kind, x in blocks(md):
        if kind == "h1":
            body.append(f"<h1>{inline_html(x)}</h1>")
        elif kind == "h2":
            body.append(f"<h2>{inline_html(x)}</h2>")
        elif kind == "hr":
            body.append("<hr>"); after_hr = True
        elif kind == "ul":
            body.append("<ul>" + "".join(f"<li>{inline_html(i)}</li>" for i in x) + "</ul>")
        elif kind == "table":
            head = "".join(f"<th>{inline_html(c)}</th>" for c in x[0])
            rows = "".join("<tr>" + "".join(f"<td>{inline_html(c)}</td>" for c in r) + "</tr>"
                           for r in x[1:])
            body.append(f"<table><thead><tr>{head}</tr></thead><tbody>{rows}</tbody></table>")
        else:
            cls = ' class="byline"' if first_p else ' class="foot"' if after_hr else \
                  ' class="note"' if x.startswith("\\*") else ""
            body.append(f"<p{cls}>{inline_html(x)}</p>")
            first_p = False
    return (f"<!doctype html><html><head><meta charset='utf-8'><title>{STEM}</title>"
            f"<style>{CSS}</style></head><body>{''.join(body)}</body></html>")


def set_borders(table):
    """Write cell borders explicitly so every viewer (not only Word) draws them."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "0"), ("w:color", "B8C2CC")):
            el.set(qn(k), v)
        borders.append(el)
    table._tbl.tblPr.append(borders)


def add_hyperlink(par, text, url, size):
    """A real, clickable Word hyperlink (python-docx has no high-level API for this)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    r_id = par.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                              is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    for tag, val in (("w:color", "1F5FA8"), ("w:u", "single"), ("w:sz", str(int(size * 2)))):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        props.append(el)
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    par._p.append(link)


def add_runs(par, text, size=10):
    for part in INLINE.split(text.replace("\\*", "\x00")):
        if not part:
            continue
        t = part.replace("\x00", "*")
        if LINK.fullmatch(part):
            add_hyperlink(par, *LINK.fullmatch(part).groups(), size)
            continue
        elif part.startswith("**"):
            r = par.add_run(t[2:-2]); r.bold = True
        elif part.startswith("`"):
            r = par.add_run(t[1:-1]); r.font.name = "Menlo"; size_ = size - 1
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = par.add_run(t[1:-1]); r.italic = True
        else:
            r = par.add_run(t)
        r.font.size = Pt(size)


def to_docx(md: str, path: Path):
    d = docx.Document()
    sec = d.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.9))
    d.styles["Normal"].font.name = "Helvetica Neue"
    d.styles["Normal"].font.size = Pt(10)
    first_p, after_hr = True, False
    for kind, x in blocks(md):
        if kind in ("h1", "h2"):
            h = d.add_heading(level=1 if kind == "h1" else 2)
            add_runs(h, x, 16 if kind == "h1" else 12)
            for r in h.runs:
                r.font.color.rgb = RGBColor(0x12, 0x32, 0x4F)
        elif kind == "hr":
            after_hr = True
        elif kind == "ul":
            for item in x:
                add_runs(d.add_paragraph(style="List Bullet"), item)
        elif kind == "table":
            t = d.add_table(rows=len(x), cols=len(x[0]))
            t.style = "Table Grid"
            set_borders(t)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, row in enumerate(x):
                for ci, cell in enumerate(row):
                    p = t.cell(ri, ci).paragraphs[0]
                    add_runs(p, f"**{cell}**" if ri == 0 and not cell.startswith("**") else cell, 9)
        else:
            small = first_p or after_hr or x.startswith("\\*")
            add_runs(d.add_paragraph(), x, 8.5 if small and not first_p else 9.5 if first_p else 10)
            first_p = False
    d.save(path)


def main():
    md = SRC.read_text()
    html_path = HERE / f"{STEM}.html"
    pdf_path = HERE / f"{STEM}.pdf"
    html_path.write_text(to_html(md))
    subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                    f"--print-to-pdf={pdf_path}", html_path.as_uri()],
                   check=True, capture_output=True, timeout=120)
    html_path.unlink()
    to_docx(md, HERE / f"{STEM}.docx")
    print(pdf_path)
    print(HERE / f"{STEM}.docx")


if __name__ == "__main__":
    main()
