"""Fill the course's Problem Statement template from docs/problem_statement.md.

    .venv/bin/python docs/build_problem_statement.py [path/to/template.docx]

Markdown tables are written as "Label — text" lines, since the template's answer
boxes are single table cells.
"""
import re
import sys
from pathlib import Path

import docx
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
TEMPLATE = Path(sys.argv[1] if len(sys.argv) > 1 else
                Path.home() / "Desktop/PE6201/Assignment/PE6201_Project_Problem_Statement_Template.docx")
OUT = HERE / "PE6201_Problem_Statement_LIU_ZEYUAN_revised.docx"


def sections(md: str) -> dict[int, list[str]]:
    out, cur = {}, None
    for line in md.splitlines():
        m = re.match(r"## (\d) ·", line)
        if m:
            cur = int(m.group(1)); out[cur] = []
        elif line.startswith("---"):
            cur = None
        elif cur is not None:
            out[cur].append(line)
    return out


def to_lines(block: list[str]) -> list[str]:
    lines, header_seen = [], False
    for raw in block:
        s = raw.strip()
        if not s:
            continue
        if s.startswith("|"):
            cells = [c.strip() for c in s.strip("|").split("|")]
            if set("".join(cells)) <= set("-: "):
                continue
            if not header_seen:          # skip the table header row
                header_seen = True
                continue
            lines.append("• " + cells[0] + " — " + " · ".join(c for c in cells[1:] if c))
            continue
        header_seen = False
        lines.append("• " + s[2:] if s.startswith("- ") else s)
    return lines


def add_runs(par, text: str):
    """Write **bold** and *italic* spans as real formatting; drop backticks."""
    text = text.replace("`", "")
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**"):
            run = par.add_run(part[2:-2]); run.bold = True
        elif part.startswith("*"):
            run = par.add_run(part[1:-1]); run.italic = True
        else:
            run = par.add_run(part)
        run.font.size = Pt(9.5)


def main():
    md = (HERE / "problem_statement.md").read_text()
    secs = sections(md)
    d = docx.Document(TEMPLATE)

    for p in d.paragraphs:
        if p.text.startswith("Name:"):
            for r in p.runs[1:]:
                r.text = ""
            p.runs[0].text = ("Name: LIU ZEYUAN     Section (A / B / C): B     "
                              "Date: 2026-09-27 (revised; original 2026-08-21)")
        if p.text.startswith("9 · Smallest first version"):
            for r in p.runs:
                r.text = r.text.replace(" (optional)", "")

    for i in range(1, 10):
        cell = d.tables[i - 1].cell(0, 0)
        for p in list(cell.paragraphs)[1:]:
            p._element.getparent().remove(p._element)
        first = cell.paragraphs[0]
        for r in first.runs:
            r.text = ""
        for k, line in enumerate(to_lines(secs[i])):
            add_runs(first if k == 0 else cell.add_paragraph(), line)

    note = md.split("---")[-1].strip().strip("*")
    add_runs(d.tables[8].cell(0, 0).add_paragraph(), f"*{note}*")
    d.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
