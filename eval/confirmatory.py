"""Confirmatory test of the flag-coverage guard. Protocol: eval/confirmatory/PREREGISTRATION.md

    .venv/bin/python -m eval.confirmatory generate   # 25 model calls (~US$0.004)
    .venv/bin/python -m eval.confirmatory sheet      # blind grading sheet (needs zh.json)
    .venv/bin/python -m eval.confirmatory analyze    # after grading

Paired design: each case gets ONE model call with prompt v2 (version A). Version B replays
exactly that model output through the frozen v3 code path, so A and B differ only by the guard.
"""
from __future__ import annotations

import json
import random
import re
import sys
from math import comb
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from coachfit import explain as ex
from coachfit.rules import assess, needs_human_review
from eval.run_eval import RESULTS, client, load_nhanes, nhanes_profile

HERE = Path(__file__).resolve().parent / "confirmatory"
SEED = 20261010
STRATA = [("below_kcal_floor", 5), ("underweight_bmi", 3), ("label_sensitive", 7),
          ("below_bmr", 5), ("none", 5)]
FLAG_COLS = ["label", "label_sensitive", "below_kcal_floor", "below_bmr", "underweight_bmi", "review"]
FLAG_ZH = {"label": "缺口标签", "label_sensitive": "误差敏感", "below_kcal_floor": "低于安全下限",
           "below_bmr": "低于 BMR", "underweight_bmi": "BMI<18.5", "review": "需要复核"}
NOTE_KEYS = ("summary", "explanation", "next_step")


# ------------------------------------------------------------------ helpers

def used_ids() -> set[str]:
    used = set()
    for f in RESULTS.rglob("nhanes_sample_ids.csv"):
        used |= {f"N{i}" for i in pd.read_csv(f).participant_id}
    for f in RESULTS.rglob("flag_fidelity_cases.json"):
        used |= {c["case_id"] for c in json.loads(f.read_text())}
    for f in RESULTS.rglob("heldout_v3.jsonl"):
        used |= {json.loads(l)["case_id"] for l in f.read_text().splitlines()}
    return used


class Replay:
    """A client that returns a stored model reply, so v3 re-uses v2's exact output."""

    def __init__(self, note: dict):
        resp = SimpleNamespace(
            choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps(note)))],
            usage=SimpleNamespace(prompt_tokens=0, completion_tokens=0))
        self.chat = SimpleNamespace(completions=SimpleNamespace(create=lambda **kw: resp))


def flags_raised(a) -> list[str]:
    out = ["label"] + [f for f in FLAG_COLS[1:5] if f in a.flags]
    if needs_human_review(a):
        out.append("review")
    return out


def detector_missing(p, a, shown: dict) -> set[str]:
    """What the frozen code considers missing from a shown note."""
    miss = {"label" if m.startswith("label:") else m for m in ex.missing_flags(p, a, shown)}
    if needs_human_review(a) and "review" not in (shown["summary"] + shown["next_step"]).lower():
        miss.add("review")
    return miss


def clopper_pearson(k: int, n: int, alpha: float = 0.05) -> tuple[float, float]:
    def cdf(x, p):
        return sum(comb(n, i) * p ** i * (1 - p) ** (n - i) for i in range(x + 1))

    def solve(f):
        lo, hi = 0.0, 1.0
        for _ in range(60):
            mid = (lo + hi) / 2
            lo, hi = (mid, hi) if f(mid) else (lo, mid)
        return (lo + hi) / 2

    lower = 0.0 if k == 0 else solve(lambda p: 1 - cdf(k - 1, p) < alpha / 2)
    upper = 1.0 if k == n else solve(lambda p: cdf(k, p) > alpha / 2)
    return lower, upper


def profiles_by_id() -> dict:
    return {f"N{r.participant_id}": nhanes_profile(r) for r in load_nhanes().itertuples()}


# ------------------------------------------------------------------ generate

def generate():
    out_path = HERE / "cases.jsonl"
    if out_path.exists():
        raise SystemExit(f"{out_path} exists: generation runs once (see PREREGISTRATION.md).")
    used = used_ids()
    pool: dict[str, list] = {}
    for cid, p in profiles_by_id().items():
        a = assess(p)
        if cid in used or a.status != "ok":
            continue
        key = next((f for f, _ in STRATA[:-1] if f in a.flags), "none")
        pool.setdefault(key, []).append(cid)
    rng = random.Random(SEED)
    picked = [(key, cid) for key, n in STRATA for cid in rng.sample(sorted(pool[key]), n)]

    llm = client()
    profiles = profiles_by_id()
    recs = []
    for stratum, cid in picked:
        p = profiles[cid]
        a = assess(p)
        A = ex.explain(p, a, client=llm, prompt="v2")
        if A["source"] == "llm":
            B = ex.explain(p, a, client=Replay(A["llm_output"]), prompt="v3")
        else:
            B = A
        recs.append({"case_id": cid, "stratum": stratum, "flags": flags_raised(a),
                     "possible_labels": a.possible_labels, "label": a.deficit_label,
                     "A": {k: A[k] for k in NOTE_KEYS}, "A_source": A["source"],
                     "B": {k: B[k] for k in NOTE_KEYS}, "B_source": B["source"],
                     "guard_added": B.get("missing_flags_raw", []) if B["source"] == "llm+guard" else [],
                     "cost_usd": A["cost_usd"]})
    out_path.write_text("\n".join(json.dumps(r) for r in recs))
    print(f"{len(recs)} cases · overlap with earlier runs: {len({r['case_id'] for r in recs} & used)}"
          f" · guard changed {sum(r['A'] != r['B'] for r in recs)} notes"
          f" · cost US${sum(r['cost_usd'] for r in recs):.4f}")


# ------------------------------------------------------------------ sheet

def sheet():
    recs = [json.loads(l) for l in (HERE / "cases.jsonl").read_text().splitlines()]
    zh = json.loads((HERE / "zh.json").read_text())       # {"<case_id>:A": ..., "<case_id>:B": ...}
    notes = [(r, v) for r in recs for v in ("A", "B")]
    random.Random(SEED).shuffle(notes)
    key = [{"note": i + 1, "case_id": r["case_id"], "version": v} for i, (r, v) in enumerate(notes)]
    (HERE / "key.json").write_text(json.dumps(key, indent=1))

    wb = openpyxl.Workbook()
    g = wb.active
    g.title = "评分说明"
    for row in [
        ("这是什么", "25 个全新案例，每个有两条说明，共 50 条，已打乱顺序，不标版本。请逐条独立评分。"),
        ("不要打开", "eval/confirmatory/key.json（版本答案）。评完之前不要看。"),
        ("逐个标记打分", "黄色格子是这条说明需要说到的标记：说到了（意思对即可）填 1，没说到填 0。灰色的「—」表示这个案例没有这个标记，不用填。"),
        ("缺口标签", "说出了结论方向即可：no_deficit=吃多了/盈余/没有缺口；too_small=缺口太小；appropriate=合适/正常；too_large=缺口太大/太激进。"),
        ("误差敏感", "提到记录误差、10%、标签可能变、核对饮食记录/记录准确性、置信度低 之一即可。"),
        ("低于安全下限", "提到低于 1200/1500 kcal 下限（floor/minimum）。"),
        ("低于 BMR", "提到摄入低于 BMR。"),
        ("BMI<18.5", "提到体重过轻 / BMI 低于 18.5。"),
        ("需要复核", "明确让教练先复核/审查这个案例。"),
        ("无编造/矛盾", "1 = 没问题；0 = 说了规则没标的标记、建议方向和标签相反（如 appropriate 却叫客户多吃）、或者前后自相矛盾。"),
        ("中文翻译", "由 Claude 直译，仅供阅读，以英文为准。"),
        ("负数缺口", "\"deficit of -90 kcal\" 等于多吃 90 kcal，不算错。too_small 时\"摄入高于目标范围\"是对的。"),
    ]:
        g.append(row)
    g.column_dimensions["A"].width = 16
    g.column_dimensions["B"].width = 110
    for row in g.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.create_sheet("评分")
    ws.append(["编号", "说明（英文原文）", "中文翻译（仅供阅读）"] + [FLAG_ZH[f] for f in FLAG_COLS]
              + ["无编造/矛盾", "你的理由"])
    yellow = PatternFill("solid", fgColor="FFF2CC")
    grey = PatternFill("solid", fgColor="E7E6E6")
    dv = DataValidation(type="list", formula1='"1,0"')
    ws.add_data_validation(dv)
    for i, (r, v) in enumerate(notes):
        n = i + 2
        cells = [i + 1, "\n".join(r[v][k] for k in NOTE_KEYS), zh[f"{r['case_id']}:{v}"]]
        cells += ["" if f in r["flags"] else "—" for f in FLAG_COLS] + ["", ""]
        ws.append(cells)
        for j, f in enumerate(FLAG_COLS):
            c = ws.cell(n, 4 + j)
            c.fill = yellow if f in r["flags"] else grey
            if f in r["flags"]:
                dv.add(c.coordinate)
        ws.cell(n, 10).fill = yellow
        dv.add(ws.cell(n, 10).coordinate)
    for col, w in zip("ABCDEFGHIJK", [6, 60, 50, 9, 9, 9, 9, 9, 9, 10, 30]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "B2"
    out = HERE / "盲评表_50条.xlsx"
    wb.save(out)
    print(out)


# ------------------------------------------------------------------ analyze

def analyze(path: str | None = None):
    recs = {r["case_id"]: r for r in map(json.loads, (HERE / "cases.jsonl").read_text().splitlines())}
    key = {k["note"]: k for k in json.loads((HERE / "key.json").read_text())}
    ws = openpyxl.load_workbook(path or HERE / "盲评表_50条_已评分.xlsx", data_only=True)["评分"]
    profiles = profiles_by_id()

    grades = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        note = row[0]
        k = key[note]
        r = recs[k["case_id"]]
        marks = {f: row[3 + j] for j, f in enumerate(FLAG_COLS) if f in r["flags"]}
        bad = [f for f, m in marks.items() if str(m).strip() not in ("0", "1")]
        if bad or str(row[9]).strip() not in ("0", "1"):
            raise SystemExit(f"note {note}: fill every yellow cell with 1 or 0")
        stated = {f: int(m) for f, m in marks.items()}
        clean = int(row[9])
        grades[(k["case_id"], k["version"])] = {
            "stated": stated, "clean": clean,
            "pass": int(all(stated.values()) and clean == 1), "reason": row[10] or ""}

    ids = sorted(recs)
    A = [grades[(c, "A")]["pass"] for c in ids]
    B = [grades[(c, "B")]["pass"] for c in ids]
    n = len(ids)

    # H2: detector recall on version A omissions
    caught = missed = false_alarm = 0
    omissions = {"A": {}, "B": {}}
    for c in ids:
        p = profiles[c]
        a = assess(p)
        det = detector_missing(p, a, recs[c]["A"])
        for f, s in grades[(c, "A")]["stated"].items():
            if s == 0:
                omissions["A"][f] = omissions["A"].get(f, 0) + 1
                caught, missed = (caught + 1, missed) if f in det else (caught, missed + 1)
            elif f in det:
                false_alarm += 1
        for f, s in grades[(c, "B")]["stated"].items():
            if s == 0:
                omissions["B"][f] = omissions["B"].get(f, 0) + 1
    total_omitted = caught + missed
    recall = caught / total_omitted if total_omitted else None
    h2_pass = (missed == 0) if total_omitted < 10 else (recall >= 0.95)
    h3 = sum(1 - grades[(c, "B")]["clean"] for c in ids)

    paired = pd.crosstab(pd.Series(A, name="A pass"), pd.Series(B, name="B pass"))
    res = {
        "n_cases": n,
        "H1_B_pass": f"{sum(B)}/{n}", "H1_B_pass_CI95": [round(x, 3) for x in clopper_pearson(sum(B), n)],
        "H1_pass": sum(B) >= 24,
        "H2_omitted_flags_in_A": total_omitted, "H2_caught": caught, "H2_missed": missed,
        "H2_recall": None if recall is None else round(recall, 3), "H2_pass": h2_pass,
        "H3_B_invented_or_contradiction": h3, "H3_pass": h3 == 0,
        "A_pass": f"{sum(A)}/{n}", "A_pass_CI95": [round(x, 3) for x in clopper_pearson(sum(A), n)],
        "A_invented_or_contradiction": sum(1 - grades[(c, "A")]["clean"] for c in ids),
        "omissions_by_flag": omissions, "detector_false_alarms_in_A": false_alarm,
        "paired_table": paired.to_dict(),
        "guard_changed_notes": sum(recs[c]["A"] != recs[c]["B"] for c in ids),
    }
    (HERE / "results.json").write_text(json.dumps(res, indent=2, default=str))
    rows = [{"case_id": c, "version": v, **grades[(c, v)]} for c in ids for v in ("A", "B")]
    pd.DataFrame(rows).to_csv(HERE / "grades.csv", index=False)
    print(json.dumps(res, indent=2, default=str))
    print("\npaired table (rows = A, columns = B):\n", paired)


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else ""
    if cmd == "generate":
        generate()
    elif cmd == "sheet":
        sheet()
    elif cmd == "analyze":
        analyze(sys.argv[2] if len(sys.argv) > 2 else None)
    else:
        raise SystemExit(__doc__)
